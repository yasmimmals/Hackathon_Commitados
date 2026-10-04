"""Carrega o cadastro de chapas (matrícula -> identificador) da planilha do boletim.

Uso (dentro do container):
    docker compose exec backend python -m scripts.carregar_chapas dados/boletim_diario_chapas.xlsx

A planilha original NÃO vai para o repositório (LEIA-ME): copie para backend/dados/,
que está no .gitignore. O cadastro fica nas colunas N (MATRICULA) e O (nome) da Plan1.
Idempotente: matrícula já cadastrada é ignorada.
"""
import sys

import openpyxl
from sqlalchemy import select

from app.core.database import SessionLocal
from app.models import Chapa


def ler_cadastro(caminho: str) -> list[tuple[str, str]]:
    ws = openpyxl.load_workbook(caminho, data_only=True).active
    cadastro = []
    for linha in range(1, ws.max_row + 1):
        matricula, nome = ws.cell(linha, 14).value, ws.cell(linha, 15).value
        if matricula is None or nome is None or str(matricula).strip().upper().startswith("MATRIC"):
            continue
        cadastro.append((str(matricula).strip(), str(nome).strip()))
    return cadastro


def run(caminho: str) -> None:
    cadastro = ler_cadastro(caminho)
    db = SessionLocal()
    try:
        existentes = set(db.scalars(select(Chapa.matricula)))
        novos = [Chapa(matricula=m, nome=n) for m, n in cadastro if m not in existentes]
        db.add_all(novos)
        db.commit()
        print(f"   chapas na planilha: {len(cadastro)} | inseridos: {len(novos)} | "
              f"já existiam: {len(cadastro) - len(novos)}")
    finally:
        db.close()


if __name__ == "__main__":
    if len(sys.argv) != 2:
        sys.exit(__doc__)
    run(sys.argv[1])