"""Boletim em Excel (.xlsx) com a formatação da Cocapec: cabeçalho, seções, moeda e totais.

Usa os mesmos valores da API (para_saida), então o arquivo bate com a tela.
"""
from decimal import Decimal
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

AZUL = "0A5AA4"
AMARELO = "F2B705"
CINZA = "F3F4F6"
MOEDA = '"R$" #,##0.00'
MOEDA_4 = '"R$" #,##0.0000'
INTEIRO = "#,##0"
DECIMAL_1 = "#,##0.0"

_fino = Side(style="thin", color="D1D5DB")
_borda = Border(left=_fino, right=_fino, top=_fino, bottom=_fino)
_titulo_secao = Font(bold=True, italic=True, size=12, color=AZUL)
_cabecalho = Font(bold=True, color="FFFFFF")
_fundo_cabecalho = PatternFill("solid", fgColor=AZUL)
_fundo_total = PatternFill("solid", fgColor=CINZA)

LOCAIS = {"INSUMOS": "Armazém de Insumos", "ADUBO": "Pátio de Adubo",
          "MAQUINAS": "Armazém de Máquinas", "LOJA": "Loja"}


def _num(v) -> float:
    return float(v) if isinstance(v, Decimal) else float(v or 0)


class _Planilha:
    """Escreve linha a linha, controlando a posição atual."""

    def __init__(self, ws):
        self.ws = ws
        self.linha = 1

    def pular(self, n: int = 1):
        self.linha += n

    def secao(self, texto: str, colunas: int):
        c = self.ws.cell(self.linha, 1, texto)
        c.font = _titulo_secao
        for col in range(1, colunas + 1):
            self.ws.cell(self.linha, col).border = Border(bottom=Side(style="medium", color=AMARELO))
        self.pular()

    def cabecalho(self, titulos: list[str]):
        for col, t in enumerate(titulos, 1):
            c = self.ws.cell(self.linha, col, t)
            c.font, c.fill, c.border = _cabecalho, _fundo_cabecalho, _borda
            c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        self.pular()

    def linha_dados(self, valores: list, formatos: list[str | None], total: bool = False):
        for col, (v, fmt) in enumerate(zip(valores, formatos), 1):
            c = self.ws.cell(self.linha, col, v)
            c.border = _borda
            if fmt:
                c.number_format = fmt
            if total:
                c.font, c.fill = Font(bold=True), _fundo_total
        self.pular()

    def par(self, rotulo: str, valor, fmt: str | None = None, destaque: bool = False):
        r = self.ws.cell(self.linha, 1, rotulo)
        r.font = Font(bold=True)
        v = self.ws.cell(self.linha, 2, valor)
        if fmt:
            v.number_format = fmt
        if destaque:
            v.font = Font(bold=True, color=AZUL, size=12)
        self.pular()


def gerar_xlsx(b: dict) -> bytes:
    """`b` é o dicionário de boletim_service.para_saida."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Boletim"
    p = _Planilha(ws)
    c = b["calculo"]
    data_br = b["data"].strftime("%d/%m/%Y")

    # Título
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=7)
    t = ws.cell(1, 1, "COCAPEC • Boletim Diário de Serviços dos Ensacadores")
    t.font = Font(bold=True, italic=True, size=14, color="FFFFFF")
    t.fill = _fundo_cabecalho
    t.alignment = Alignment(vertical="center")
    ws.row_dimensions[1].height = 26
    p.pular(2)

    p.par("Data", data_br)
    p.par("Armazém", LOCAIS.get(getattr(b["local"], "value", b["local"]) or "", "Geral (Franca)"))
    status = getattr(b["status"], "value", b["status"])
    p.par("Status", "Fechado" if status == "FECHADO" else "Rascunho")
    if b.get("fechado_por"):
        fechado_em = b["fechado_em"].strftime("%d/%m/%Y %H:%M") if b.get("fechado_em") else ""
        p.par("Fechado por", f"{b['fechado_por']} {fechado_em}".strip())
    if b.get("observacao"):
        p.par("Observação", b["observacao"])
    p.pular()

    # Produção
    p.secao("1. Produção do dia", 7)
    p.cabecalho(["Tipo de item", "Descarga", "Remoção", "Transferência", "Quantidade", "Preço unitário", "Valor"])
    for l in b["linhas"]:
        p.linha_dados(
            [l["tipo_item"], l["descarga"], l["remocao"], l["transferencia"], l["quantidade_total"],
             _num(l["preco_unitario"]), _num(l["valor_linha"])],
            [None, INTEIRO, INTEIRO, INTEIRO, INTEIRO, MOEDA_4, MOEDA])
    if not b["linhas"]:
        p.linha_dados(["Sem produção lançada", "", "", "", "", "", 0], [None] * 6 + [MOEDA])
    p.linha_dados(["Total da produção", "", "", "", sum(l["quantidade_total"] for l in b["linhas"]), "",
                   _num(c["producao_total"])], [None, None, None, None, INTEIRO, None, MOEDA], total=True)
    p.pular()

    # Equipe
    p.secao("2. Equipe temporária", 6)
    p.cabecalho(["Nº", "Matrícula", "Nome", "Diária", "Diária equivalente", "Pagamento"])
    por_diaria = max(_num(c["valor_por_diaria"]), _num(c["piso_diaria"]))
    for i, ch in enumerate(b["equipe"], 1):
        peso = 0.5 if ch["meia_diaria"] else 1.0
        p.linha_dados([i, ch["matricula"], ch["nome"], "Meia" if ch["meia_diaria"] else "Completa", peso,
                       por_diaria * peso], [INTEIRO, None, None, None, DECIMAL_1, MOEDA])
    p.linha_dados(["", "", f"{c['chapas']} chapa(s)", f"{c['meias_diarias']} meia(s)",
                   _num(c["diarias_equivalentes"]), _num(c["total_pagar"])],
                  [None, None, None, None, DECIMAL_1, MOEDA], total=True)
    p.pular()

    # Fechamento
    p.secao("3. Fechamento", 2)
    p.par("Produção total", _num(c["producao_total"]), MOEDA)
    p.par("Diárias equivalentes", _num(c["diarias_equivalentes"]), DECIMAL_1)
    p.par("Valor por diária", _num(c["valor_por_diaria"]) if c["valor_por_diaria"] is not None else "—", MOEDA)
    p.par("Piso por diária", _num(c["piso_diaria"]), MOEDA_4)
    p.par("Piso total da equipe", _num(c["piso_total"]), MOEDA)
    p.par("Complemento pago pela cooperativa", _num(c["complemento"]), MOEDA)
    p.par("Total a pagar", _num(c["total_pagar"]), MOEDA, destaque=True)

    if b["avisos"]:
        p.pular()
        p.secao("Avisos", 7)
        for a in b["avisos"]:
            ws.cell(p.linha, 1, a).font = Font(color="92400E")
            p.pular()

    for col, largura in enumerate([34, 14, 30, 14, 18, 16, 16], 1):
        ws.column_dimensions[get_column_letter(col)].width = largura
    ws.sheet_view.showGridLines = False
    ws.print_options.horizontalCentered = True
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1

    saida = BytesIO()
    wb.save(saida)
    return saida.getvalue()
