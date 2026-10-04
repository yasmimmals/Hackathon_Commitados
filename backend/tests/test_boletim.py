"""Boletim diário (Tarefa 2). Os números conferem com o exemplo da seção 8 do dossiê."""
from datetime import date
from decimal import Decimal as D

import openpyxl
import pytest
from fastapi.testclient import TestClient

from app.core.config import PISO_DIARIA
from app.core.database import SessionLocal
from app.models import BoletimDiario, Chapa, TipoItem
from app.services.boletim_calculo import arredondar, calcular
from scripts.carregar_chapas import ler_cadastro

PRECO = D("0.3224")

# Cadastro real da planilha (colunas N/O): matrícula -> identificador anonimizado
CHAPAS_REAIS = [("158", "CHAPA_08"), ("137", "CHAPA_09"), ("35", "CHAPA_15"),
                ("155", "CHAPA_48"), ("69", "CHAPA_37"), ("13", "CHAPA_38"),
                ("75", "CHAPA_41"), ("74", "CHAPA_42"), ("79", "CHAPA_43"),
                ("15", "CHAPA_49"), ("161", "CHAPA_30"), ("1", "CHAPA_24")]
EQUIPE_DO_EXEMPLO = ["158", "137", "35", "155", "161", "15", "69", "13", "75", "74", "79"]


def r(v):
    return arredondar(v)


# ------------------------------------------------------------ cálculo puro

def test_t1_exemplo_do_dossie():
    """Adubo, 17/11/2025: 2.848 itens a 0,3224, 11 chapas em diária completa."""
    res = calcular([(2848, PRECO)], n_chapas=11, n_meias=0, piso=PISO_DIARIA)
    assert r(res.producao_total) == D("918.20")
    assert res.diarias_equivalentes == 11
    assert r(res.valor_por_diaria) == D("83.47")
    assert r(res.total_pagar) == D("991.90")
    assert r(res.complemento) == D("73.71")
    assert res.abaixo_do_piso


def test_t2_um_chapa_em_meia_diaria():
    res = calcular([(2848, PRECO)], n_chapas=11, n_meias=1, piso=PISO_DIARIA)
    assert res.diarias_equivalentes == D("10.5")
    assert r(res.valor_por_diaria) == D("87.45")
    assert r(res.total_pagar) == D("946.82")
    assert r(res.complemento) == D("28.62")


def test_t3_producao_acima_do_piso_nao_tem_complemento_nem_teto():
    res = calcular([(5000, PRECO)], n_chapas=10, n_meias=0, piso=PISO_DIARIA)
    assert r(res.producao_total) == D("1612.00")              # 161,20 por diária > piso
    assert res.complemento == 0 and res.total_pagar == res.producao_total


def test_armadilha_de_arredondar_no_meio():
    """Por que o banco não guarda 2 casas: arredondando antes de subtrair, dá 73,70."""
    res = calcular([(2848, PRECO)], 11, 0, PISO_DIARIA)
    errado = r(res.total_pagar) - r(res.producao_total)
    assert errado == D("73.70") and r(res.complemento) == D("73.71")


def test_sem_equipe_nao_divide_por_zero():
    res = calcular([(100, PRECO)], 0, 0, PISO_DIARIA)
    assert res.valor_por_diaria is None and res.complemento == 0
    assert res.total_pagar == res.producao_total


def test_meia_diaria_vale_meio_piso_e_nao_45_0786():
    """A planilha mostra 'Meia diária R$ 45,0786', mas a fórmula usa 0,5 × 90,1731."""
    res = calcular([], n_chapas=1, n_meias=1, piso=PISO_DIARIA)
    assert res.total_pagar == D("45.08655")


def test_mais_meias_que_chapas_e_invalido():
    with pytest.raises(ValueError):
        calcular([], 1, 2, PISO_DIARIA)


# ------------------------------------------------------------ API

@pytest.fixture
def chapas(db):
    db.add_all([Chapa(matricula=m, nome=n) for m, n in CHAPAS_REAIS])
    db.commit()


@pytest.fixture
def api():
    from main import app
    return TestClient(app)


@pytest.fixture
def tipo(db):
    def _tipo(descricao):
        return db.query(TipoItem).filter_by(descricao=descricao).one().id
    return _tipo


def producao_do_exemplo(tipo):
    return {"linhas": [
        {"tipo_item_id": tipo("Fertilizantes"), "descarga": 2378, "remocao": 400},
        {"tipo_item_id": tipo("Agroquímico"), "descarga": 30},
        {"tipo_item_id": tipo("Serviços diversos"), "remocao": 40},
    ]}


def test_fluxo_completo_do_exemplo_pela_api(api, chapas, tipo):
    b = api.post("/api/v1/boletins", json={"data": "2025-11-17"}).json()
    assert b["status"] == "RASCUNHO" and b["local"] is None      # boletim geral do dia

    b = api.put(f"/api/v1/boletins/{b['id']}/producao", json=producao_do_exemplo(tipo)).json()
    fert = next(ln for ln in b["linhas"] if ln["tipo_item"] == "Fertilizantes")
    assert fert["quantidade_total"] == 2778 and fert["valor_linha"] == "895.63"
    assert b["calculo"]["producao_total"] == "918.20"

    equipe = {"equipe": [{"matricula": m} for m in EQUIPE_DO_EXEMPLO]}
    b = api.put(f"/api/v1/boletins/{b['id']}/equipe", json=equipe).json()
    c = b["calculo"]
    assert (c["chapas"], c["diarias_equivalentes"]) == (11, "11.0")
    assert (c["valor_por_diaria"], c["total_pagar"], c["complemento"]) == ("83.47", "991.90", "73.71")
    assert c["abaixo_do_piso"] and any("complemento" in a for a in b["avisos"])
    assert {e["nome"] for e in b["equipe"]} >= {"CHAPA_08", "CHAPA_43"}

    # T2: um deles em meia diária
    equipe["equipe"][0]["meia_diaria"] = True
    c = api.put(f"/api/v1/boletins/{b['id']}/equipe", json=equipe).json()["calculo"]
    assert (c["diarias_equivalentes"], c["total_pagar"], c["complemento"]) == ("10.5", "946.82", "28.62")

    b = api.post(f"/api/v1/boletins/{b['id']}/fechar", json={"fechado_por": "resp.adubo"}).json()
    assert b["status"] == "FECHADO" and b["fechado_por"] == "resp.adubo"

    # o valor gravado é o exato (946,81755 - 918,1952), sem arredondamento no meio
    with SessionLocal() as s:
        gravado = s.get(BoletimDiario, b["id"])
        assert gravado.complemento == D("28.62235") and r(gravado.complemento) == D("28.62")

    r1 = api.put(f"/api/v1/boletins/{b['id']}/equipe", json=equipe)
    assert r1.status_code == 409 and r1.json()["codigo"] == "BOLETIM_FECHADO"
    assert api.post(f"/api/v1/boletins/{b['id']}/reabrir").json()["status"] == "RASCUNHO"


def test_mesma_matricula_duas_vezes_e_recusada(api, chapas):
    """Sem esta trava, o mesmo chapa contaria 3 diárias e inflaria o custo."""
    b = api.post("/api/v1/boletins", json={"data": "2025-11-18"}).json()
    r1 = api.put(f"/api/v1/boletins/{b['id']}/equipe",
                 json={"equipe": [{"matricula": "158"}] * 3})
    assert r1.status_code == 409 and r1.json()["codigo"] == "MATRICULA_REPETIDA"


def test_matricula_nao_cadastrada(api, chapas):
    b = api.post("/api/v1/boletins", json={"data": "2025-11-18"}).json()
    r1 = api.put(f"/api/v1/boletins/{b['id']}/equipe",
                 json={"equipe": [{"matricula": "158"}, {"matricula": "999"}]})
    assert r1.json()["codigo"] == "CHAPA_NAO_CADASTRADA"
    assert r1.json()["detalhes"]["matriculas"] == ["999"]


def test_limite_de_20_chapas(api, db):
    db.add_all([Chapa(matricula=str(1000 + i), nome=f"T{i}") for i in range(21)])
    db.commit()
    b = api.post("/api/v1/boletins", json={"data": "2025-11-18"}).json()
    r1 = api.put(f"/api/v1/boletins/{b['id']}/equipe",
                 json={"equipe": [{"matricula": str(1000 + i)} for i in range(21)]})
    assert r1.json()["codigo"] == "LIMITE_CHAPAS"


def test_um_boletim_geral_por_dia(api):
    primeiro = api.post("/api/v1/boletins", json={"data": "2025-11-18"}).json()
    assert primeiro["local"] is None
    r1 = api.post("/api/v1/boletins", json={"data": "2025-11-18"})
    assert r1.status_code == 409 and r1.json()["codigo"] == "BOLETIM_JA_EXISTE"
    assert r1.json()["detalhes"]["boletim_id"] == primeiro["id"]
    assert api.post("/api/v1/boletins", json={"data": "2025-11-19"}).status_code == 201


def test_boletim_geral_nao_aceita_armazem(api):
    r1 = api.post("/api/v1/boletins", json={"data": "2025-11-18", "local": "ADUBO"})
    assert r1.status_code == 409 and r1.json()["codigo"] == "BOLETIM_GERAL"


# ------------------------------------------------------------ modo por armazém (dossiê)

@pytest.fixture
def por_armazem(monkeypatch):
    monkeypatch.setattr("app.core.config.BOLETIM_POR_ARMAZEM", True)


def test_modo_por_armazem_um_por_armazem_por_dia(api, por_armazem):
    corpo = {"data": "2025-11-18", "local": "INSUMOS"}
    assert api.post("/api/v1/boletins", json=corpo).status_code == 201
    r1 = api.post("/api/v1/boletins", json=corpo)
    assert r1.json()["codigo"] == "BOLETIM_JA_EXISTE"
    assert api.post("/api/v1/boletins", json={**corpo, "local": "ADUBO"}).status_code == 201
    r2 = api.post("/api/v1/boletins", json={"data": "2025-11-18"})
    assert r2.json()["codigo"] == "LOCAL_OBRIGATORIO"


def test_geral_e_por_armazem_nunca_no_mesmo_dia(api, monkeypatch):
    """Senão o custo do dia seria contado duas vezes."""
    api.post("/api/v1/boletins", json={"data": "2025-11-18"})
    monkeypatch.setattr("app.core.config.BOLETIM_POR_ARMAZEM", True)
    r1 = api.post("/api/v1/boletins", json={"data": "2025-11-18", "local": "ADUBO"})
    assert r1.status_code == 409 and r1.json()["codigo"] == "BOLETIM_JA_EXISTE"


def test_tipo_de_item_repetido(api, tipo):
    b = api.post("/api/v1/boletins", json={"data": "2025-11-18"}).json()
    linha = {"tipo_item_id": tipo("Fertilizantes"), "descarga": 10}
    r1 = api.put(f"/api/v1/boletins/{b['id']}/producao", json={"linhas": [linha, linha]})
    assert r1.json()["codigo"] == "TIPO_ITEM_REPETIDO"


def test_nao_fecha_sem_equipe(api):
    b = api.post("/api/v1/boletins", json={"data": "2025-11-18"}).json()
    r1 = api.post(f"/api/v1/boletins/{b['id']}/fechar", json={"fechado_por": "x"})
    assert r1.json()["codigo"] == "BOLETIM_SEM_EQUIPE"


def test_aviso_de_chapa_em_dois_armazens_no_mesmo_dia(api, chapas, por_armazem):
    """Modo por armazém: não bloqueia (pode ter trabalhado nos dois), mas avisa."""
    for local in ("ADUBO", "INSUMOS"):
        b = api.post("/api/v1/boletins", json={"data": "2025-11-19", "local": local}).json()
        b = api.put(f"/api/v1/boletins/{b['id']}/equipe",
                    json={"equipe": [{"matricula": "158"}]}).json()
    assert any("158" in a and "ADUBO" in a for a in b["avisos"])


def test_preco_fica_congelado_no_lancamento(api, db, tipo):
    b = api.post("/api/v1/boletins", json={"data": "2025-11-20"}).json()
    api.put(f"/api/v1/boletins/{b['id']}/producao",
            json={"linhas": [{"tipo_item_id": tipo("Fertilizantes"), "descarga": 1000}]})
    t = db.get(TipoItem, tipo("Fertilizantes"))
    t.preco_unitario = D("0.5000")                     # reajuste depois do lançamento
    db.commit()
    assert api.get(f"/api/v1/boletins/{b['id']}").json()["calculo"]["producao_total"] == "322.40"


def test_listagem_e_tipos_item(api):
    api.post("/api/v1/boletins", json={"data": "2025-11-17"})
    api.post("/api/v1/boletins", json={"data": "2025-11-25"})
    lista = api.get("/api/v1/boletins", params={"inicio": "2025-11-01", "fim": "2025-11-20"}).json()
    assert [b["data"] for b in lista] == ["2025-11-17"]
    tipos = api.get("/api/v1/cadastros/tipos-item").json()
    assert len(tipos) == 14 and tipos[0]["preco_unitario"] == "0.1824"


# ------------------------------------------------------------ carga do cadastro

def test_carregar_chapas_da_planilha(tmp_path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.cell(1, 14, "MATRICULA "), ws.cell(1, 15, "FUNCIONARIOS  ADUBO")
    for i, (m, n) in enumerate(CHAPAS_REAIS[:3], start=3):
        ws.cell(i, 14, int(m)), ws.cell(i, 15, n)
    caminho = tmp_path / "boletim.xlsx"
    wb.save(caminho)
    assert ler_cadastro(str(caminho)) == CHAPAS_REAIS[:3]


def test_data_do_boletim_aceita_sabado(api):
    """Aos sábados a equipe faz organização interna: o boletim registra toda movimentação."""
    assert date(2025, 11, 22).weekday() == 5
    assert api.post("/api/v1/boletins", json={"data": "2025-11-22"}).status_code == 201